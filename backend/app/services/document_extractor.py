import io
import re
from pathlib import Path
from pypdf import PdfReader
import docx
from app.config import ALLOWED_EXTENSIONS

def clean_extracted_text(text: str) -> str:
    text = re.sub(r'\n\s*\n', '\n\n', text)
    text = re.sub(r'[ \t]+', ' ', text)
    return text.strip()

def extract_text_from_pdf(file_stream: io.BytesIO) -> str:
    reader = PdfReader(file_stream)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text

def extract_text_from_docx(file_stream: io.BytesIO) -> str:
    doc = docx.Document(file_stream)
    text_parts = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text.strip())

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
            if row_text:
                text_parts.append(row_text)

    return "\n".join(text_parts)

def extract_text_from_txt(file_bytes: bytes) -> str:
    try:
        return file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        return file_bytes.decode("latin-1", errors="ignore")

def extract_document_text(file_bytes: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()

    if ext not in ALLOWED_EXTENSIONS:
        allowed_str = ", ".join(sorted(ALLOWED_EXTENSIONS))
        raise ValueError(f"File format '{ext}' is not supported. Allowed formats: {allowed_str}")

    file_stream = io.BytesIO(file_bytes)
    raw_text = ""

    if ext == ".pdf":
        raw_text = extract_text_from_pdf(file_stream)
    elif ext == ".docx":
        raw_text = extract_text_from_docx(file_stream)
    elif ext == ".txt":
        raw_text = extract_text_from_txt(file_bytes)

    cleaned = clean_extracted_text(raw_text)

    if not cleaned:
        raise ValueError("Failed to extract text from file. Ensure the document contains selectable text and is not an image scan.")

    return cleaned
